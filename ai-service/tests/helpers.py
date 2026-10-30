"""Build small in-memory documents for tests."""

from io import BytesIO

from docx import Document
from pypdf import PdfWriter


def build_text_pdf(text: str) -> bytes:
    """Create a tiny valid PDF containing one line of Latin-1 text."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")

    xref_position = len(output)
    total = len(objects) + 1
    output.extend(f"xref\n0 {total}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer << /Size {total} /Root 1 0 R >>\n"
            f"startxref\n{xref_position}\n%%EOF"
        ).encode("ascii")
    )
    return bytes(output)


def build_blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_encrypted_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        column_count = len(table_rows[0])
        table = document.add_table(rows=len(table_rows), cols=column_count)
        for row_index, row_values in enumerate(table_rows):
            for column_index, value in enumerate(row_values):
                table.cell(row_index, column_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
