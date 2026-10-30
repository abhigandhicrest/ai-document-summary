"""Extract text from PDF, DOCX, and TXT uploads.

Uploaded files are treated as data. This module never executes them.
"""

import re
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import EmptyFileError, FileNotDecryptedError, PdfReadError

from app.errors import AppServiceError

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt"}


def extract_text(filename: str, data: bytes) -> str:
    """Identify a supported document and return its normalized text."""
    if data is None or len(data) == 0:
        raise AppServiceError("EMPTY_FILE", "The uploaded file is empty.", 400)

    file_type = _identify_format(filename, data)
    if file_type == "pdf":
        text = _read_pdf(data)
        empty_message = (
            "No usable text could be extracted from this PDF. "
            "Scanned PDFs need OCR, which is not supported in this version."
        )
    elif file_type == "docx":
        text = _read_docx(data)
        empty_message = "No usable text could be extracted from this Word document."
    else:
        text = _read_txt(data)
        empty_message = "This text file has no usable content."

    normalized = normalize_text(text)
    if not normalized:
        raise AppServiceError("NO_EXTRACTABLE_TEXT", empty_message, 422)
    return normalized


def normalize_text(text: str) -> str:
    """Collapse extra spaces while keeping paragraph breaks."""
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0c", "\n")
    paragraphs = []
    for block in re.split(r"\n\s*\n", cleaned):
        lines = [" ".join(line.split()) for line in block.split("\n")]
        paragraph = " ".join(line for line in lines if line).strip()
        if paragraph:
            paragraphs.append(paragraph)
    return "\n\n".join(paragraphs)


def public_filename(filename: str) -> str:
    """Return a display name that cannot be used as a filesystem path."""
    base_name = Path(filename or "document").name
    cleaned = re.sub(r"[^\w.\- ()]", "_", base_name).strip("._ ")
    return (cleaned or "document")[:180]


def _identify_format(filename: str, data: bytes) -> str:
    extension = Path(Path(filename or "").name).suffix.lower()
    file_type = SUPPORTED_EXTENSIONS.get(extension)
    if file_type is None:
        raise AppServiceError(
            "UNSUPPORTED_FILE_TYPE",
            "Upload a PDF, DOCX, or TXT file.",
            415,
        )

    if file_type == "pdf" and not _looks_like_pdf(data):
        raise AppServiceError(
            "INVALID_DOCUMENT",
            "The file extension says PDF, but the contents are not a PDF.",
            400,
        )
    if file_type == "docx" and not _looks_like_docx(data):
        raise AppServiceError(
            "INVALID_DOCUMENT",
            "The file extension says DOCX, but the contents are not a Word document.",
            400,
        )
    if file_type == "txt" and b"\x00" in data[:8192]:
        raise AppServiceError(
            "INVALID_DOCUMENT",
            "This text file contains binary data and cannot be read.",
            400,
        )
    return file_type


def _looks_like_pdf(data: bytes) -> bool:
    return b"%PDF" in data[:1024]


def _looks_like_docx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _read_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            # Some PDFs are marked encrypted but open with an empty password.
            decrypted = reader.decrypt("")
            if decrypted == 0:
                raise AppServiceError(
                    "ENCRYPTED_PDF",
                    "This PDF is encrypted. Remove the password and upload it again.",
                    422,
                )
    except AppServiceError:
        raise
    except EmptyFileError as error:
        raise AppServiceError("EMPTY_FILE", "The uploaded file is empty.", 400) from error
    except FileNotDecryptedError as error:
        raise AppServiceError(
            "ENCRYPTED_PDF",
            "This PDF is encrypted. Remove the password and upload it again.",
            422,
        ) from error
    except PdfReadError as error:
        raise AppServiceError(
            "CORRUPTED_PDF",
            "This PDF could not be read. The file may be corrupted.",
            422,
        ) from error
    except Exception as error:
        raise AppServiceError(
            "CORRUPTED_PDF",
            "This PDF could not be read. The file may be corrupted.",
            422,
        ) from error

    page_text = []
    try:
        for page in reader.pages:
            extracted = page.extract_text() or ""
            if extracted.strip():
                page_text.append(extracted)
    except PdfReadError as error:
        raise AppServiceError(
            "CORRUPTED_PDF",
            "This PDF could not be read. The file may be corrupted.",
            422,
        ) from error
    return "\n\n".join(page_text)


def _read_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as error:
        raise AppServiceError(
            "CORRUPTED_DOCX",
            "This Word document could not be read. The file may be corrupted.",
            422,
        ) from error

    parts = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = _unique_cell_text(row.cells)
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _unique_cell_text(cells) -> list[str]:
    values = []
    for cell in cells:
        value = " ".join(cell.text.split())
        if value and (not values or values[-1] != value):
            values.append(value)
    return values


def _read_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise AppServiceError(
            "ENCODING_ERROR",
            "This text file is not valid UTF-8. Save it as UTF-8 and try again.",
            422,
        ) from error
